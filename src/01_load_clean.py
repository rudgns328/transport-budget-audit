"""01. 지방재정365 세부사업별 세출현황 XLSX 적재·정제.

입력: data/raw/{연도}_세부사업별세출현황.xlsx
출력:
  data/processed/budget_{연도}.parquet   정제된 세출 (0원 행 제거)
  outputs/qa/anomalies_{연도}.csv        점검용 이상치 목록 (행을 지우지 않음)
  outputs/qa/01_load_clean_report.md     보고 항목 요약

실행: python src/01_load_clean.py [--years 2025 2024 2019]
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
PROCESSED_DIR = ROOT / "data" / "processed"
QA_DIR = ROOT / "outputs" / "qa"

# 2단 머리글(2~3행)의 기대값. 원본 구조가 바뀌면 바로 멈추도록 검사한다.
EXPECTED_HEADER_ROW2 = ("지역", "자치단체", "회계", "세부사업명", "예산현액", None, None, None, None,
                        "지출액", "집행잔액", "분야", "부문")
EXPECTED_HEADER_ROW3 = (None, None, None, None, "계", "국비", "시도비", "시군구비", "기타",
                        None, None, None, None)
COLUMNS = ["지역", "자치단체", "회계", "세부사업명",
           "예산현액", "예산_국비", "예산_시도비", "예산_시군구비", "예산_기타",
           "지출액", "집행잔액", "분야", "부문"]
AMOUNT_COLS = ["예산현액", "예산_국비", "예산_시도비", "예산_시군구비", "예산_기타", "지출액", "집행잔액"]
TEXT_COLS = ["지역", "자치단체", "회계", "세부사업명", "분야", "부문"]
DATA_START_ROW = 4  # 엑셀 행 번호 기준


def raw_path(year: int) -> Path:
    return RAW_DIR / f"{year}_세부사업별세출현황.xlsx"


def read_header(path: Path) -> tuple[str, tuple, tuple]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(max_row=3, values_only=True))
    wb.close()
    return rows[0][0], tuple(rows[1]), tuple(rows[2])


def normalize_text(s):
    """줄바꿈·탭·연속 공백을 공백 하나로, 앞뒤 공백 제거."""
    if pd.isna(s):
        return s
    return re.sub(r"\s+", " ", str(s)).strip()


def load(year: int) -> pd.DataFrame:
    path = raw_path(year)
    unit, h2, h3 = read_header(path)
    if unit != "단위:원" or h2 != EXPECTED_HEADER_ROW2 or h3 != EXPECTED_HEADER_ROW3:
        raise ValueError(f"{path.name}: 머리글 구조가 예상과 다릅니다\n{unit}\n{h2}\n{h3}")

    df = pd.read_excel(path, header=None, skiprows=DATA_START_ROW - 1, names=COLUMNS, dtype=object)
    df.insert(0, "id", [f"{year}-{i + DATA_START_ROW}" for i in range(len(df))])  # 엑셀 행 번호로 원본 추적
    df.insert(1, "연도", year)
    return df


def clean(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    stats = {"원본 행 수": len(df)}

    df["세부사업명_원문"] = df["세부사업명"]
    for c in TEXT_COLS:
        df[c] = df[c].map(normalize_text)
    stats["사업명 공백·줄바꿈 정리된 행"] = int(
        (df["세부사업명_원문"].astype(str) != df["세부사업명"].astype(str)).sum())
    stats["사업명 빈 행"] = int(df["세부사업명"].isna().sum() + (df["세부사업명"] == "").sum())

    for c in AMOUNT_COLS:
        num = pd.to_numeric(df[c], errors="coerce")
        stats[f"숫자 변환 실패({c})"] = int((num.isna() & df[c].notna()).sum())
        stats[f"결측({c})"] = int(df[c].isna().sum())
        df[c] = num.fillna(0).astype("int64")

    zero = (df["예산현액"] == 0) & (df["지출액"] == 0)
    stats["0원 행 수(예산현액=0, 지출액=0)"] = int(zero.sum())
    stats["0원 행 비율"] = round(zero.mean() * 100, 2)
    df = df.loc[~zero].reset_index(drop=True)
    stats["정제 후 행 수"] = len(df)

    df["광역기초"] = df["자치단체"].eq("본청").map({True: "광역", False: "기초"})
    return df, stats


def find_anomalies(df: pd.DataFrame) -> pd.DataFrame:
    """점검용 이상치. 행을 제거하지 않고 목록만 만든다."""
    checks = {
        "지출액>예산현액": df["지출액"] > df["예산현액"],
        "음수 금액": (df[AMOUNT_COLS] < 0).any(axis=1),
        "예산현액≠재원 합": df["예산현액"] != df[["예산_국비", "예산_시도비", "예산_시군구비", "예산_기타"]].sum(axis=1),
        "집행잔액≠예산현액-지출액": df["집행잔액"] != df["예산현액"] - df["지출액"],
    }
    flags = pd.DataFrame(checks)
    out = df.loc[flags.any(axis=1), ["id", "지역", "자치단체", "회계", "세부사업명", "분야", "부문"] + AMOUNT_COLS].copy()
    out.insert(1, "이상 유형", flags.loc[flags.any(axis=1)].apply(lambda r: "; ".join(r.index[r]), axis=1))
    return out


def md_table(s: pd.Series, name: str, value: str = "행 수") -> str:
    lines = [f"| {name} | {value} |", "| --- | ---: |"]
    lines += [f"| {k} | {v:,} |" for k, v in s.items()]
    return "\n".join(lines)


def report_year(year: int, df: pd.DataFrame, stats: dict, anomalies: pd.DataFrame) -> str:
    out = [f"## {year}년", ""]
    out += [f"- {k}: {v:,}" if isinstance(v, int) else f"- {k}: {v}" for k, v in stats.items()]
    n_gov = df[["지역", "자치단체"]].drop_duplicates()
    out += [
        f"- 자치단체 수 (지역, 자치단체 조합): {len(n_gov):,} (광역 본청 {(n_gov['자치단체'] == '본청').sum()}, "
        f"기초 {(n_gov['자치단체'] != '본청').sum()})",
        f"- 고유 세부사업명 수: {df['세부사업명'].nunique():,}",
        f"- 지출액 합계: {df['지출액'].sum():,} 원 / 예산현액 합계: {df['예산현액'].sum():,} 원",
        f"- 동일 (지역, 자치단체, 회계, 사업명) 중복 행: {df.duplicated(['지역', '자치단체', '회계', '세부사업명']).sum():,}",
        "",
        "### 이상치 (anomalies 파일 참고, 행은 제거하지 않음)", "",
        md_table(anomalies["이상 유형"].str.split("; ").explode().value_counts(), "이상 유형"),
        "",
        "### 지역별 자치단체 수", "",
        md_table(n_gov.groupby("지역", sort=False).size(), "지역", "자치단체 수"),
        "",
        "### 회계 종류", "",
        md_table(df["회계"].value_counts(), "회계"),
        "",
        "### 분야", "",
        md_table(df["분야"].value_counts(dropna=False), "분야"),
        "",
        "### 분야·부문", "",
        md_table(df.groupby(["분야", "부문"], dropna=False).size().rename(lambda x: x).pipe(
            lambda s: s.set_axis([f"{a} / {b}" for a, b in s.index])), "분야 / 부문"),
        "",
    ]
    return "\n".join(out)


def region_checks(frames: dict[int, pd.DataFrame]) -> str:
    """3-3 지역 매칭 규칙 관련 표기 확인."""
    out = ["## 지역 표기 확인 (3-3)", ""]
    for year, df in frames.items():
        gov = df[["지역", "자치단체"]].drop_duplicates()
        out.append(f"- {year} 지역 표기: {', '.join(gov['지역'].unique())}")
        for name in ["군위군"]:
            hit = gov[gov["자치단체"] == name]
            out.append(f"- {year} {name}: {', '.join(hit['지역']) or '없음'}")
        for region in ["세종", "제주"]:
            hit = gov[gov["지역"].str.startswith(region)]
            out.append(f"- {year} {region} 자치단체: {', '.join(hit['자치단체']) or '없음'}")
        dup = gov[gov.duplicated("자치단체", keep=False) & (gov["자치단체"] != "본청")]
        out.append(f"- {year} 여러 시도에 있는 같은 이름: "
                   + ", ".join(f"{k}({v})" for k, v in dup.groupby("자치단체")["지역"].agg(", ".join).items()))
    out.append("")
    return "\n".join(out)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--years", type=int, nargs="+", default=[2025, 2024, 2019])
    args = ap.parse_args()

    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)

    sections, frames, headers = [], {}, {}
    for year in args.years:
        print(f"[{year}] 적재 중: {raw_path(year).name}")
        headers[year] = read_header(raw_path(year))
        df, stats = clean(load(year))
        anomalies = find_anomalies(df)
        df.to_parquet(PROCESSED_DIR / f"budget_{year}.parquet", index=False)
        anomalies.to_csv(QA_DIR / f"anomalies_{year}.csv", index=False, encoding="utf-8-sig")
        frames[year] = df
        sections.append(report_year(year, df, stats, anomalies))
        print(f"[{year}] 정제 후 {len(df):,}행 저장, 이상치 {len(anomalies):,}행")

    head = ["# 01 적재·정제 보고", "", "자동 생성: src/01_load_clean.py", ""]
    if len(headers) > 1:
        same = len({h for h in headers.values()}) == 1
        head += [f"- 머리글·열 구성 연도 간 동일: {'예' if same else '아니오'} ({', '.join(map(str, headers))})", ""]
    report = "\n".join(head + [region_checks(frames)] + sections)
    (QA_DIR / "01_load_clean_report.md").write_text(report, encoding="utf-8")
    print(f"보고서: {QA_DIR / '01_load_clean_report.md'}")


if __name__ == "__main__":
    main()

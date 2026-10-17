"""03. 사람 라벨링용 층화 표본 200건과 빈 양식 생성.

- LLM 결과를 쓰지 않는다. 표본 단위는 후보 행(id), 모집단은 3개년 후보 전체
- 층: 추출단계 × 부문. 2단계에서 기대 배정이 1.5건 미만인 부문은 '2단계|소규모 부문 통합'으로 묶는다
- 1단계 층은 최소 2건 배정, 나머지는 행 수 비례 (최대 잉여 방식). 05에서 층 비중으로 가중 정확도 계산
- 프롬프트 [예시] 사업명과 같은 사업은 제외 (○○ 자리는 임의 문자열로 보고 패턴 매칭)
- 같은 사업명(공백 무시)이 두 번 뽑히지 않게 한다

입력: data/processed/candidates.parquet, prompts/classify_v1_1.txt, CLAUDE.md(4장)
출력:
  data/labels/label_template.xlsx   시트1 '라벨링'(빈 라벨 열), 시트2 '기준'(CLAUDE.md 4장 그대로)
  data/labels/sample_ids.csv        번호 ↔ id, 층 정보, 층 모집단 행 수 (04·05에서 사용)

실행: python src/03_sample_for_labeling.py [--force]
라벨이 입력됐을 수 있는 기존 양식은 --force 없이는 덮어쓰지 않는다.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
CANDIDATES = ROOT / "data" / "processed" / "candidates.parquet"
PROMPT = ROOT / "prompts" / "classify_v1_1.txt"
CLAUDE_MD = ROOT / "CLAUDE.md"
LABEL_DIR = ROOT / "data" / "labels"

SEED = 20261017
N_SAMPLE = 200
MIN_STAGE1 = 2
STAGE2_MIN_EXPECTED = 1.5
CATEGORIES = list("ABCDEFGH")


def name_key(s: str) -> str:
    return re.sub(r"\s+", "", s)


def example_patterns() -> list[re.Pattern]:
    """프롬프트 [예시] 블록의 큰따옴표 안 사업명을 정규식으로 만든다."""
    text = PROMPT.read_text(encoding="utf-8")
    block = text.split("[예시]", 1)[1].split("[출력 형식]", 1)[0]
    names = re.findall(r'"([^"]+)"', block)
    pats = []
    for n in names:
        parts = [re.escape(name_key(p)) for p in n.split("○○")]
        pats.append(re.compile("^" + ".+".join(parts) + "$"))
    return pats


def assign_strata(df: pd.DataFrame) -> pd.Series:
    stratum = df["추출단계"].astype(str) + "단계|" + df["부문"]
    s2 = df["추출단계"] == 2
    expected = stratum[s2].value_counts() / len(df) * N_SAMPLE
    small = expected[expected < STAGE2_MIN_EXPECTED].index
    stratum[s2 & stratum.isin(small)] = "2단계|소규모 부문 통합"
    return stratum


def allocate(sizes: pd.Series) -> pd.Series:
    """1단계 층 최소 MIN_STAGE1건, 나머지는 비례 배정 (최대 잉여)."""
    floor = pd.Series(0, index=sizes.index)
    floor[sizes.index.str.startswith("1단계")] = MIN_STAGE1
    floor = np.minimum(floor, sizes)
    rest = N_SAMPLE - floor.sum()
    quota = sizes / sizes.sum() * rest
    alloc = floor + np.floor(quota).astype(int)
    remainder = (quota - np.floor(quota)).sort_values(ascending=False)
    for k in remainder.index[: N_SAMPLE - alloc.sum()]:
        alloc[k] += 1
    return np.minimum(alloc, sizes).astype(int)


def draw(df: pd.DataFrame, alloc: pd.Series, rng: np.random.Generator) -> pd.DataFrame:
    picked, used_keys = [], set()
    for stratum, n in alloc.items():
        pool = df[df["층"] == stratum]
        order = rng.permutation(len(pool))
        got = 0
        for i in order:
            row = pool.iloc[i]
            if row["사업명키"] in used_keys:
                continue
            picked.append(row)
            used_keys.add(row["사업명키"])
            got += 1
            if got == n:
                break
        if got < n:
            raise RuntimeError(f"{stratum}: 고유 사업명이 부족합니다 ({got}/{n})")
    return pd.DataFrame(picked)


def criteria_section() -> list[str]:
    text = CLAUDE_MD.read_text(encoding="utf-8")
    m = re.search(r"^## 4\..*?(?=^## 5\.)", text, re.S | re.M)
    if not m:
        raise ValueError("CLAUDE.md에서 4장을 찾지 못했습니다")
    return m.group(0).rstrip().splitlines()


def write_template(sample: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "라벨링"
    ws.append(["입력 안내: 범주는 A~H 중 하나, 확신도는 high 또는 low. "
               "판단 기준은 '기준' 시트만 사용하고 다른 사람·AI 결과는 보지 마세요."])
    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(bold=True, color="C00000")
    header = ["번호", "사업명", "분야", "부문", "범주", "확신도", "메모"]
    ws.append(header)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for _, r in sample.iterrows():
        ws.append([int(r["번호"]), r["세부사업명"], r["분야"], r["부문"], None, None, None])
    last = ws.max_row
    for col, choices in (("E", CATEGORIES), ("F", ["high", "low"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True,
                            showErrorMessage=True, error=f"{', '.join(choices)} 중 하나를 입력하세요")
        dv.add(f"{col}3:{col}{last}")
        ws.add_data_validation(dv)
    for col, w in zip("ABCDEFG", [6, 60, 18, 22, 8, 9, 40]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"

    ws2 = wb.create_sheet("기준")
    ws2.append(["CLAUDE.md 4장을 그대로 옮김 (분류 기준 v1.1)"])
    ws2["A1"].font = Font(bold=True)
    for line in criteria_section():
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(set(c) <= {"-", ":", " "} for c in cells):
                continue
            ws2.append(cells)
        else:
            ws2.append([line])
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 30
    ws2.column_dimensions["C"].width = 80
    for row in ws2.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=False, vertical="top")
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="기존 양식을 덮어쓴다 (라벨 입력 전에만)")
    args = ap.parse_args()

    template = LABEL_DIR / "label_template.xlsx"
    ids_path = LABEL_DIR / "sample_ids.csv"
    if (template.exists() or ids_path.exists()) and not args.force:
        raise SystemExit(f"{template.name} 또는 {ids_path.name}이 이미 있습니다. 라벨 입력 전이면 --force로 다시 만드세요.")

    df = pd.read_parquet(CANDIDATES)
    df["사업명키"] = df["세부사업명"].map(name_key)
    pats = example_patterns()
    is_example = df["사업명키"].map(lambda k: any(p.match(k) for p in pats))
    n_excluded = int(is_example.sum())
    df = df[~is_example].reset_index(drop=True)

    df["층"] = assign_strata(df)
    sizes = df["층"].value_counts()
    alloc = allocate(sizes)
    sample = draw(df, alloc[alloc > 0], np.random.default_rng(SEED))
    sample = sample.sample(frac=1, random_state=SEED).reset_index(drop=True)  # 층별로 몰리지 않게 섞는다
    sample["번호"] = np.arange(1, len(sample) + 1)
    sample["층 모집단 행 수"] = sample["층"].map(sizes)
    sample["층 표본 수"] = sample["층"].map(alloc)

    LABEL_DIR.mkdir(parents=True, exist_ok=True)
    cols = ["번호", "id", "연도", "지역", "자치단체", "세부사업명", "분야", "부문", "추출단계", "층",
            "층 모집단 행 수", "층 표본 수"]
    sample[cols].to_csv(ids_path, index=False, encoding="utf-8-sig")
    write_template(sample, template)

    print(f"프롬프트 예시와 같은 사업 제외: {n_excluded}행 (패턴 {len(pats)}개)")
    print(f"표본 {len(sample)}건 → {template.relative_to(ROOT)}, {ids_path.relative_to(ROOT)}")
    summary = pd.DataFrame({"모집단 행 수": sizes, "표본": alloc}).sort_index()
    print(summary.to_string())
    print("연도별 표본:", sample["연도"].value_counts().sort_index(ascending=False).to_dict())
    print("광역/기초:", sample["자치단체"].eq("본청").map({True: "광역", False: "기초"}).value_counts().to_dict())


if __name__ == "__main__":
    main()
